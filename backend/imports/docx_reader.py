"""
Word DOCX Table Reader with Embedded Cell Photo Extraction

Features:
- Parses .docx Word document tables into structured student records.
- Scans table cells for drawing elements (w:drawing, a:blip), extracts relationship
  image binaries directly from the docx package, and assigns them to card photo fields.
"""
import io
import os
import logging
from typing import Dict, Any, List, Tuple, Optional
from dataclasses import dataclass, field

from .excel_reader import ExtractedImage, ParsedSpreadsheet

logger = logging.getLogger(__name__)


def parse_docx_tables(file_bytes: bytes) -> ParsedSpreadsheet:
    """
    Parse a Word (.docx) document, locate its primary data table,
    and extract text cells + embedded cell photos.
    """
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError(f"Could not read Word .docx document: {exc}")

    if not doc.tables:
        raise ValueError("No data tables found in this Word document. Please ensure your data is inside a table.")

    # Find the largest table (most rows)
    primary_table = max(doc.tables, key=lambda t: len(t.rows))

    if len(primary_table.rows) < 2:
        raise ValueError("The table in the Word document must contain at least a header row and one data row.")

    # 1. Parse Header Row
    header_cells = primary_table.rows[0].cells
    headers = [cell.text.strip() for cell in header_cells]
    # Filter empty trailing headers
    while headers and not headers[-1]:
        headers.pop()

    if not headers:
        headers = [f"FIELD_{i+1}" for i in range(len(header_cells))]

    num_cols = len(headers)
    rows = []
    embedded_images: List[ExtractedImage] = []

    # Helper namespace mappings for finding blip elements
    blip_xpath = './/a:blip'
    embed_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'

    # 2. Parse Data Rows
    for data_row_idx, row in enumerate(primary_table.rows[1:]):
        row_values = []
        for c_idx in range(num_cols):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell_text = cell.text.strip()
                row_values.append(cell_text)

                # Scan cell XML for embedded drawings/photos
                try:
                    for blip in cell._element.xpath(blip_xpath):
                        rId = blip.get(embed_attr)
                        if rId and rId in doc.part.related_parts:
                            image_part = doc.part.related_parts[rId]
                            image_bytes = image_part.blob
                            if image_bytes and len(image_bytes) > 100:
                                ext = os.path.splitext(image_part.partname)[1].lstrip('.').lower() or 'jpeg'
                                field_name = headers[c_idx] if c_idx < len(headers) else f"FIELD_{c_idx+1}"

                                embedded_images.append(ExtractedImage(
                                    row_idx=data_row_idx,
                                    col_idx=c_idx,
                                    field_name=field_name,
                                    image_bytes=image_bytes,
                                    format=ext,
                                ))
                except Exception as cell_img_err:
                    logger.debug("Failed extracting docx cell image at row %d col %d: %s", data_row_idx, c_idx, cell_img_err)
            else:
                row_values.append('')

        if any(v != '' for v in row_values) or any(img.row_idx == data_row_idx for img in embedded_images):
            rows.append(row_values)

    return ParsedSpreadsheet(
        headers=headers,
        rows=rows,
        embedded_images=embedded_images,
        total_rows=len(rows),
        total_embedded_images=len(embedded_images),
    )
